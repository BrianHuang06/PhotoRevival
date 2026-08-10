# -*- coding: utf-8 -*-
from pathlib import Path

import docx
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

doc = Document()

style = doc.styles['Normal']
style.font.name = 'SimSun'
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimSun')
style.paragraph_format.line_spacing = 1.5

for level in range(1, 4):
    hs = doc.styles[f'Heading {level}']
    hs.font.name = 'SimHei'
    hs.font.color.rgb = RGBColor(0, 0, 0)
    hs.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')
    if level == 1:
        hs.font.size = Pt(18)
    elif level == 2:
        hs.font.size = Pt(15)
    else:
        hs.font.size = Pt(13)

def add_para(text, bold=False, size=None, align=None, space_after=Pt(6)):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if size:
        run.font.size = size
    if align:
        p.alignment = align
    p.paragraph_format.space_after = space_after
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = h
        for par in cell.paragraphs:
            for run in par.runs:
                run.bold = True
                run.font.size = Pt(11)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = table.rows[ri + 1].cells[ci]
            cell.text = str(val)
            for par in cell.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(11)
    doc.add_paragraph()
    return table


# ===== 封面 =====
for _ in range(6):
    doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('尘拂光回')
run.font.size = Pt(36)
run.bold = True
run.font.name = 'SimHei'
run.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('基于SDXL与LoRA微调的老照片智能修复系统')
run.font.size = Pt(18)
run.font.name = 'SimHei'
run.element.rPr.rFonts.set(qn('w:eastAsia'), 'SimHei')

doc.add_paragraph()

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.add_run('中国国际大学生创新大赛 · 创意项目赛道')
run.font.size = Pt(14)

for _ in range(4):
    doc.add_paragraph()

info_lines = [
    '项目负责人：黄仕昌',
    '指导教师：王超群',
    '项目归属学院：人工智能学院',
    '华南师范大学',
]
for line in info_lines:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(line)
    run.font.size = Pt(14)

doc.add_page_break()

# ===== 一、项目概述 =====
doc.add_heading('一、项目概述', level=1)

add_para(
    '"尘拂光回"是一个基于深度学习的老照片智能修复系统。项目创新性地将"损伤检测—退化去除—细节重建"三阶段串联，构建了两阶段解耦修复架构：第一阶段采用SwinIR Transformer执行退化去除，消除模糊、噪声与压缩伪影并实现4倍超分辨率增强；第二阶段基于Stable Diffusion XL Inpainting模型配合LoRA领域微调权重，在损伤掩码引导下进行精准的细节重建。'
)

add_para(
    '项目自主实现了基于Gabor滤波+LBP+随机森林的多特征融合划痕检测引擎（5折交叉验证准确率达98.6%），自主设计了退化合成数据管线，低成本构建了涵盖划痕、污渍、褪色、暗角、颗粒噪声等多种退化效果的配对训练数据。基于开源SDXL Inpainting预训练模型和SwinIR图像恢复模型，通过LoRA参数高效微调进行领域适配，经过多版本迭代优化（v1至v29），逐步筛选最优模型配置。'
)

add_para(
    '项目已完成从模型训练到前端系统的全栈开发，测试结果显示SNR改善率达94.6%，对比度改善率约30%。系统可广泛应用于博物馆档案数字化保护、家庭影像修复、影视后期制作等领域，为"AI赋能文化遗产保护"提供了一条可复用的技术路径与实践范式。'
)

# ===== 二、需求分析与项目背景 =====
doc.add_heading('二、需求分析与项目背景', level=1)

doc.add_heading('（一）老照片正在消失', level=2)
add_para(
    '老照片是不可再生的历史记忆载体。我国馆藏老照片超过2000万张，约70%存在划痕、污渍、褪色、撕裂等不同程度的损伤。这些照片记录着国家发展、城市变迁、家族记忆，一旦损毁便无法复原。随着时间推移，照片材质持续老化，修复窗口正在迅速缩小。'
)

doc.add_heading('（二）现有修复手段的困境', level=2)
add_para(
    '人工修复依赖专业技师逐像素修补，单张耗时2-5天，费用200-800元，全国具备此能力的技师不足千人，供需严重失衡。市面上的AI修复工具（如Remini、美图秀秀等）采用通用图像增强模型，对老照片的复合退化（划痕+褪色+模糊+噪声）缺乏针对性处理，修复结果普遍存在三大问题：过度平滑——面部五官、衣物纹理等细节被模糊处理；色彩失真——修复后的色调偏离照片的年代特征；内容篡改——AI会"脑补"出原始照片中不存在的内容，破坏照片的历史真实性。Adobe Photoshop等专业软件虽具备修复能力，但需手动标注修复区域、逐层操作，使用门槛极高，单张处理耗时30-60分钟，无法满足批量修复需求。'
)

doc.add_heading('（三）技术契机', level=2)
add_para(
    '近年来以Stable Diffusion为代表的扩散模型在图像生成与修复领域取得突破性进展，为老照片修复提供了全新的技术路径。然而，通用扩散模型直接应用于老照片修复时存在"过度生成"问题——修复结果虽然清晰，但往往偏离原始照片的真实内容。如何让AI在去除损伤的同时"尊重"照片的原始面貌，成为该领域尚未解决的核心挑战。本项目正是针对这一技术空白提出的创意方案。'
)

doc.add_heading('（四）政策支撑', level=2)
add_para(
    '《关于加强文物数字化工作的指导意见》（文物科函〔2023〕123号）明确提出"推动文物影像资料的数字化采集、修复与保护"；《"十四五"文化发展规划》将"文化遗产数字化保护"列为重点工程；广东省《关于推进文化强省建设的实施意见》提出"加快文化资源的数字化转化与开发利用"。本项目与上述政策方向高度契合。'
)

# ===== 三、技术方案 =====
doc.add_heading('三、技术方案', level=1)

doc.add_heading('（一）整体架构', level=2)
add_para(
    '系统采用"检测—退化去除—细节重建"三阶段串联架构。核心创新在于将修复过程解耦为两个独立阶段，每个阶段聚焦单一目标，避免"既要去除损伤又要保留细节"的多目标冲突。'
)
add_para('完整技术流程：输入老照片 → [Stage 0] 多特征融合损伤检测，生成损伤掩码(Mask) → [Stage 1] SwinIR退化去除，去模糊/去噪/去压缩伪影/4倍超分 → [Stage 2] SDXL Inpainting + LoRA细节重建，掩码引导定向修复 → 输出修复结果。', bold=True)

doc.add_heading('（二）损伤检测引擎', level=2)
add_para(
    '本项目自主研发了基于多特征融合的损伤检测方案，区别于传统的形态学操作（如黑帽变换），采用机器学习方法实现更精准的损伤定位。'
)
add_para(
    '特征提取方面，设计了72个不同方向与频率的Gabor滤波核（8个方向 × 3个sigma × 3个lambda），捕获多尺度的划痕纹理特征；同时提取LBP（局部二值模式）直方图区分划痕与正常纹理，结合Sobel梯度方向直方图识别线性损伤特征，并加入局部统计特征（均值、方差、偏度、峰度、熵）和频率特征。五类特征拼接后形成416维的高维描述向量。'
)
add_para(
    '模型训练方面，从250张Grunge纹理素材中提取3000个划痕正样本，从146张干净照片中提取2920个正常负样本，训练200棵决策树的随机森林分类器。经5折交叉验证，分类准确率达到98.6%（±1.68%）。'
)
add_para(
    '推理流程采用stride=4的滑窗策略逐块分类，生成像素级损伤概率热图，经阈值化与形态学后处理输出最终二值掩码。'
)

doc.add_heading('（三）第一阶段：SwinIR退化去除', level=2)
add_para(
    '采用SwinIR（Swin Image Restoration）模型，基于Swin Transformer的移位窗口自注意力机制，对输入图像执行全局退化去除。该阶段负责消除模糊、噪声、压缩伪影等均匀性退化，并实现4倍超分辨率增强。其输出作为第二阶段扩散模型的条件输入，确保训练与推理阶段条件输入的一致性。'
)

doc.add_heading('（四）第二阶段：SDXL + LoRA细节重建', level=2)
add_para(
    '这是整个系统的技术核心。基于Stable Diffusion XL Inpainting模型，在损伤掩码引导下对退化区域进行定向修复。'
)

add_para('LoRA领域微调', bold=True)
add_para(
    '通用SDXL模型在老照片修复场景中表现不佳，容易生成不自然的内容。本项目采用LoRA（Low-Rank Adaptation）技术，仅训练极少量参数（占原模型0.1%以下）即可将通用模型适配至老照片修复场景。LoRA适配器插入UNet的Cross-Attention层，通过低秩矩阵分解（r=16）捕获老照片的年代风格特征——如早期相纸的暖黄基调、胶片颗粒分布规律等——在去除残留损伤的同时恢复面部、纹理等高频细节。'
)

add_para('训练数据构建', bold=True)
add_para(
    '自研退化合成管线（age_photo.py），利用250余张Grunge纹理素材，通过随机组合划痕叠加、色彩偏移、暗角模拟、颗粒噪声、褪色变色、压缩伪影等6类退化效果，低成本生成高质量配对训练数据（退化图-干净图），解决了老照片修复领域配对数据稀缺的瓶颈问题。'
)

add_para('迭代优化过程', bold=True)
add_para(
    '项目经历了从v1到v29共29轮迭代优化。每轮迭代中，通过调整LoRA秩维度(r)、扩展系数(α)、学习率、损失函数组合等超参数，在验证集上对比PSNR与SSIM指标，逐步筛选最优配置。最终选定的checkpoint在损伤去除率与细节保真度之间取得了最佳平衡。'
)

add_para('掩码引导修复策略', bold=True)
add_para(
    '在推理阶段，损伤检测引擎生成的掩码引导SDXL Inpainting模型仅对破损区域施加较高修复力度，对完好区域施加较低修复力度，实现"破损处充分修复、完好处保持原貌"的精准定位修复。'
)

doc.add_heading('（五）原型系统', level=2)
add_para(
    '项目已完成可运行的全栈原型系统。后端基于Flask REST API，前端基于React + TypeScript + Tailwind CSS，通过异步任务队列实现修复任务的实时进度反馈。系统支持三种修复模式（快速修复、精细修复、仅增强）和完整的参数调节面板，用户通过浏览器即可完成从上传到下载的全流程操作。'
)

add_para('系统核心功能：', bold=True)
add_table(
    ['功能模块', '功能描述'],
    [
        ['图片上传', '支持拖拽上传JPG/PNG格式，自动识别分辨率与文件信息'],
        ['损伤检测', '自动识别划痕、污渍等损伤区域，生成可视化掩码叠加预览'],
        ['三种修复模式', '快速修复（一键处理）、精细修复（自定义参数）、仅增强（超分辨率）'],
        ['参数调节', '修复强度、引导系数、推理步数、随机种子、混合因子可调'],
        ['修复对比', 'Before/After交互式滑动对比，支持逐像素查看'],
        ['结果导出', '下载修复结果、损伤掩码、中间过程图'],
        ['历史管理', '修复任务列表、缩略图预览、批量操作'],
    ]
)

# ===== 四、系统实现与效果展示 =====
doc.add_heading('四、系统实现与效果展示', level=1)

doc.add_heading('（一）量化评估', level=2)
add_para(
    '为全面、客观评估模型与系统的修复效果，项目构建了客观数据量化和主观感知评价的双轨评估体系。'
)
add_para('客观量化指标', bold=True)
add_para(
    '在测试集上对比优化前后模型的PSNR、SSIM等指标。目前实验结果显示：在10张测试图片上，PSNR从20.83提升至22.00，人脸区域PSNR为21.59。另对90张真实破损照片进行修复评估，SNR改善率达94.6%，对比度改善率约30%。同时也发现修复后图像锐度有所下降（去噪过程的副作用），后续需进一步优化。'
)
add_table(
    ['评估指标', '数值'],
    [
        ['PSNR提升', '20.83 → 22.00 dB'],
        ['人脸区域PSNR', '21.59 dB'],
        ['SNR改善率', '94.6%（90张真实照片）'],
        ['对比度改善率', '约30%'],
        ['划痕检测器分类准确率', '98.6%（5折交叉验证）'],
    ]
)

add_para('主观感知评分', bold=True)
add_para(
    '邀请相关专业人员及普通用户组成评价团队，从色彩真实性、面部自然度、质感还原度三个维度进行10分制打分，计算平均分与标准差，评估修复结果的主观接受度。评估实施流程为：第一阶段在测试集上自动计算定量指标，形成基础评估报告；第二阶段选取典型案例提交评估团队深度点评，根据反馈调整LoRA训练参数与损失函数权重。'
)

doc.add_heading('（二）效果展示', level=2)
add_para('（此处应插入3-5组修复前后对比图，建议包含以下案例：）', size=Pt(10))
add_para(
    '案例一：严重划痕损伤照片。经损伤检测生成掩码后，两阶段修复管线精准去除划痕，面部五官与背景纹理完整保留。'
)
add_para(
    '案例二：褪色+污渍+模糊复合损伤照片。精细修复模式下，色彩得到还原，污渍被清除，图像清晰度显著提升。'
)
add_para(
    '案例三：低分辨率压缩伪影照片。经SwinIR 4倍超分辨率增强后，文字标识与边缘细节显著恢复。'
)

# ===== 五、项目创新点 =====
doc.add_heading('五、项目创新点', level=1)

add_para('创新一：两阶段解耦修复架构', bold=True)
add_para(
    '现有方案采用单一模型端到端修复，损伤去除与细节恢复相互冲突，容易过度平滑或产生伪影。本项目创新性地将修复解耦为"退化去除"（SwinIR）与"细节重建"（SDXL+LoRA）两个独立阶段，每个子任务聚焦单一目标。这一设计使得修复效果的精细度与可控性显著提升，用户可通过参数调节在"保守修复"与"深度修复"之间灵活切换。'
)

add_para('创新二：面向老照片的LoRA领域微调', bold=True)
add_para(
    '首次将LoRA参数高效微调技术系统性地应用于老照片修复场景。通过自研退化合成数据管线构建训练数据，在SDXL Inpainting模型的UNet注意力层插入低秩适配器，以极少参数量（占原模型0.1%以下）实现通用扩散模型向老照片修复领域的精准迁移。经过29轮迭代优化，验证了LoRA在图像修复领域的有效性与工程可行性。'
)

add_para('创新三：多特征融合损伤检测与掩码引导精准修复', bold=True)
add_para(
    '自主研发了Gabor+LBP+梯度特征融合的损伤检测引擎，结合掩码引导的定向修复策略，实现了"只修复该修的地方"的精准修复理念。区别于全图无差别处理的传统方案，大幅降低了过度修改风险，保障了修复结果的历史真实性。'
)

add_para('创新四：合成退化数据管线', bold=True)
add_para(
    '自研age_photo退化合成管线，支持6类退化效果的随机组合生成，以极低成本构建高质量配对训练数据，解决了老照片修复领域"没有配对数据就无法训练"的根本性瓶颈，为同类研究提供了可复用的数据构建方案。'
)

# ===== 六、社会价值与教育意义 =====
doc.add_heading('六、社会价值与教育意义', level=1)

doc.add_heading('（一）文化遗产保护价值', level=2)
add_para(
    '老照片是不可再生的历史记忆载体。本项目为博物馆、档案馆等文化机构提供了低成本、高效率的老照片数字化修复工具，使大量因修复能力不足而持续退化的珍贵影像资料得以抢救性保护。修复后的高质量数字影像可用于历史研究、文化传播、数字展览等多个场景，持续释放文化价值。每一张被修复的老照片，都是一段被挽救的历史记忆。'
)

doc.add_heading('（二）技术普惠价值', level=2)
add_para(
    '项目将原本只有专业技师才能完成的老照片修复能力，以AI工具的形式开放给普通大众。用户无需任何专业技能，上传照片即可获得专业级的修复效果，极大地降低了文化遗产保护的参与门槛，体现了技术向善的价值取向。'
)

doc.add_heading('（三）双创教育价值', level=2)
add_para(
    '本项目是华南师范大学人工智能学院"专创融合"教育的实践成果。项目从专业课（深度学习、计算机视觉）中萌生创意，在创新创业课程的方法论指导下完成从创意到原型的完整转化。团队成员在项目实践中系统掌握了扩散模型原理、LoRA微调技术、全栈系统开发等前沿技能，实现了从"学知识"到"用知识解决真问题"的能力跃迁。学校在项目培育过程中提供了办公场地、GPU算力资源和指导教师对接等全方位支持。'
)

doc.add_heading('（四）新文科融合价值', level=2)
add_para(
    '项目将人工智能（理科）与文化遗产保护（文科）深度融合，团队在推进过程中不仅需要掌握深度学习等硬核技术，还需了解照片修复的历史保护原则、色彩学基础、档案管理规范等跨学科知识，是"技术赋能人文、人文引领技术"的典型实践，为学校新文科建设提供了可展示的成果。'
)

# ===== 七、团队介绍 =====
doc.add_heading('七、团队介绍', level=1)

add_table(
    ['成员', '专业', '职责', '相关能力'],
    [
        ['黄仕昌', '人工智能', '项目负责人、核心算法开发', 'PyTorch、扩散模型、LoRA微调、全栈开发'],
        ['成员A', '（待填写）', '（待填写）', '（待填写）'],
        ['成员B', '（待填写）', '（待填写）', '（待填写）'],
        ['成员C', '（待填写）', '（待填写）', '（待填写）'],
    ]
)

add_para('指导教师：王超群，华南师范大学人工智能学院（此处补充教师研究方向与职称）。在本项目中提供深度学习模型架构设计、训练策略优化等方面的技术指导。')

# ===== 八、未来展望 =====
doc.add_heading('八、未来展望', level=1)

add_para('技术优化方向', bold=True)
add_para(
    '对模型进行知识蒸馏与INT8量化，将显存需求从8GB降至4GB，覆盖更多消费级设备；集成GFPGAN人脸增强模块，提升面部修复精度；增加批量处理与视频修复能力，拓展应用场景。'
)

add_para('应用拓展方向', bold=True)
add_para(
    '从老照片修复拓展至油画修复、壁画修复、胶片修复等多介质文化遗产修复场景；探索与博物馆、档案馆的深度合作，构建"AI+文化遗产保护"的行业解决方案；计划开展社区公益修复活动，为老兵、退休教师等群体提供免费修复服务。'
)

add_para('开源与学术贡献', bold=True)
add_para(
    '计划将核心推理代码与合成数据管线在GitHub开源，为学术社区提供可复用的基线方案；整理项目成果撰写学术论文，推动老照片修复领域的技术进步。'
)

# ===== 保存 =====
output_path = Path(__file__).resolve().parent / 'reports' / '尘拂光回_创意赛道项目计划书.docx'
output_path.parent.mkdir(parents=True, exist_ok=True)
doc.save(output_path)
print(f'Document saved to: {output_path}')
